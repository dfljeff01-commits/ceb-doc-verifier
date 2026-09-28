# -*- coding: utf-8 -*-
"""运单核对核心逻辑（车底号/铅封号）——Issue #4 任务二。

业务场景（Issue #4 任务书）：
  A同事在发站现场装车后形成准确清单（车种/车底号/箱号/铅封号，比对基准）；
  B同事把车底号录入中欧班列订车系统后，把系统生成的加密PDF运单批量导出上传；
  本模块在"提交订车系统之后、确认运单之前"的窗口内，以箱号为主键核对两边记录，
  标出车底号/铅封号不一致的箱子，供B同事在确认运单前当场修改（错了只能回车站
  重新盖章，改不成发到边境有退运风险）。

提取实现（确定性规则：无token消耗、敏感数据不出本机；复用已验证的
"pikepdf空密码解密 + pdfplumber解析"基础，原生集成、不走技能包安装）：
  - 解密：pikepdf 空密码解密订车系统导出的加密运单；本就未加密的文件原样通过。
  - 箱号：复用既有 `[A-Z]{4}\\d{7}` 正则。按任务书仅作两边匹配主键，
    不做真伪核查（箱号源自装箱照片与货代提箱确认，信任度高）。
  - 车底号：SMGS第7栏"车辆—Вагон"。按栏位标签锚定+表格框线求解取值区域；
    必须排除第8栏"车辆由何方提供—Вагон предоставлен"的同词干扰
    （该栏也含"Вагон"，只按"含Вагон"匹配会误抓发站代码）。
  - 铅封号：SMGS第19栏"封印—Пломбы"下"记号—знаки"子栏。
    【关键约束，不可回退】必须按PDF视觉布局的坐标/框位定位取值，
    禁止按文本流顺序取"第几个数字"：版面上散布大量6位数字（车站代码、
    运单号等），且同一票货不同阶段的导出框内值不同——真实案例：
    260921导出框内是装车旧封 261918，260925导出框内是发站换封后的 260281
    （与现场清单序号48行一致）。顺序推断法有真实概率抓错成旧封。

健壮性（任务书"防止硬编码规则静默失效"，与提取同等重要）：
  - 锚点词找不到           → anchor_missing（模板可能变化）
  - 有框线但框内无数字     → not_found
  - 提取值格式异常         → format_anomaly（位数不对/含非数字），绝不静默输出
  - 框内多个互异候选       → ambiguous（不猜，交人工）
  - 逐份异常之外，批级统计异常占比，超过阈值（默认20%，可配）时给出
    "运单模板可能已变化，建议人工抽查"的醒目告警，不能静默继续输出。
"""

from __future__ import annotations

import io
import os
import re
from dataclasses import dataclass, field

# ---------------------------------------------------------------------------
# 常量与配置
# ---------------------------------------------------------------------------

#: 箱号：与全系统一致的既有口径（仅作匹配主键，不做真伪核查）
CONTAINER_RE = re.compile(r"[A-Z]{4}\d{7}")

#: 车底号默认格式：7位（现场清单口径，如5480825），兼容俄铁8位车号
DEFAULT_WAGON_RE = r"\d{7,8}"
#: 铅封号默认格式：6位纯数字（真实运单口径，如260281）
DEFAULT_SEAL_RE = r"\d{6}"

#: 批级异常占比告警阈值（任务书建议20%，可用环境变量覆盖）
DEFAULT_ANOMALY_RATIO = 0.2
ENV_ANOMALY_RATIO = "CEBS_WAYBILL_ANOMALY_RATIO"

# 栏位锚点词（订车系统现行为中文+俄文混排版式；纯俄文退化锚点兜底）
WAGON_ANCHOR_EXACT = "车辆—Вагон"          # 第7栏标签（精确整词）
WAGON_ANCHOR_RU = "Вагон"                  # 纯俄文退化：独立"Вагон"单词
WAGON_FORBID_NEIGHBOR = "предоставлен"     # 第8栏"Вагон предоставлен"排除词
SEAL_ANCHOR_EXACT = "记号—знаки"           # 第19栏"记号—знаки"子栏标签
SEAL_ANCHOR_RU = "знаки"                   # 纯俄文退化：独立"знаки"单词

# 取值区域求解参数（点，pdfplumber坐标，top向下增长）
WAGON_BAND_DEPTH = 85    # 第7栏为多行框（虚线分行的车板槽位），纵深取整框
WAGON_BAND_LEFT = 25     # 锚点向左覆盖栏位号"7"
WAGON_BAND_RIGHT = 85    # 锚点向右到第7/8栏分界线
SEAL_BAND_DEPTH = 45     # 第19栏"记号"子栏通常1-2行（多枚封印）
SEAL_BAND_LEFT = 45      # 锚点向左越过"знаки"标签左缘到子栏分界线
SEAL_BAND_RIGHT = 35     # 锚点向右到栏右边界

# 字段提取状态
FIELD_OK = "ok"
FIELD_NOT_FOUND = "not_found"              # 锚点正常但框内无候选值
FIELD_ANCHOR_MISSING = "anchor_missing"    # 锚点词找不到（模板可能变化）
FIELD_FORMAT_ANOMALY = "format_anomaly"    # 框内有值但格式不符（位数/非数字）
FIELD_AMBIGUOUS = "ambiguous"              # 框内多个互异候选，无法确定
FIELD_READ_ERROR = "read_error"            # 文件损坏/加密不可解，整体无法解析

_ANOMALY_STATUSES = (FIELD_NOT_FOUND, FIELD_ANCHOR_MISSING,
                     FIELD_FORMAT_ANOMALY, FIELD_AMBIGUOUS, FIELD_READ_ERROR)

# 批量核对状态
ST_MATCH = "match"                          # 两项一致，确认正确
ST_WAGON_MISMATCH = "wagon_mismatch"        # 车底号不一致
ST_SEAL_MISMATCH = "seal_mismatch"          # 铅封号不一致
ST_BOTH_MISMATCH = "both_mismatch"          # 两项均不一致
ST_PDF_UNMATCHED = "pdf_unmatched"          # PDF箱号在清单中找不到
ST_EXCEL_UNMATCHED = "excel_unmatched"      # 清单箱号没有对应PDF
ST_EXTRACTION_ANOMALY = "extraction_anomaly"  # 提取异常，无法判定
ST_DATA_MISSING = "data_missing"            # 清单侧字段为空，无法比较


def default_anomaly_ratio() -> float:
    """批级告警阈值：环境变量优先，默认20%（任务书建议值）。"""
    raw = os.environ.get(ENV_ANOMALY_RATIO, "").strip()
    try:
        return float(raw) if raw else DEFAULT_ANOMALY_RATIO
    except ValueError:
        return DEFAULT_ANOMALY_RATIO


# ---------------------------------------------------------------------------
# 数据结构
# ---------------------------------------------------------------------------

@dataclass
class FieldExtraction:
    """单字段的提取结果与过程留痕（锚点/候选/说明，供人工核查与回归分析）。"""
    field: str
    value: str | None = None
    status: str = FIELD_NOT_FOUND
    candidates: list[str] = field(default_factory=list)
    anchor: str | None = None        # 命中的锚点词（模板漂移排查依据）
    method: str = "coordinate"       # coordinate=框线/坐标定位；text=正则全文
    note: str = ""

    @property
    def is_anomaly(self) -> bool:
        return self.status in _ANOMALY_STATUSES


@dataclass
class WaybillExtraction:
    """一份PDF运单的完整提取结果。"""
    filename: str
    container: FieldExtraction
    wagon_no: FieldExtraction
    seal_no: FieldExtraction
    page_count: int = 1

    def anomaly_fields(self) -> list[str]:
        return [f.field for f in self.fields() if f.is_anomaly]

    def fields(self) -> list[FieldExtraction]:
        return [self.container, self.wagon_no, self.seal_no]


@dataclass
class PackingRow:
    """A同事现场装车清单的一行（事实来源/比对基准）。"""
    row_no: int                       # Excel实际行号
    seq: str | None = None            # 清单"序号"列（任务书"第48行"即此口径）
    wagon_type: str | None = None     # 车种
    wagon_no: str | None = None       # 车底号
    container_no: str | None = None   # 箱号（匹配主键）
    seal_no: str | None = None        # 铅封号


# ---------------------------------------------------------------------------
# PDF打开与解密（pikepdf空密码 + pdfplumber）
# ---------------------------------------------------------------------------

def open_pdf(data: bytes):
    """pdfplumber打开运单PDF；打不开时先做pikepdf空密码解密再试。

    订车系统导出的运单通常带加密标志；真实样本验证过空密码可解。
    两次都失败时抛出原始异常（调用方按"文件不可读"计异常，不静默跳过）。
    """
    import pdfplumber

    try:
        return pdfplumber.open(io.BytesIO(data))
    except Exception:
        decrypted = _decrypt_empty_password(data)
        return pdfplumber.open(io.BytesIO(decrypted))


def _decrypt_empty_password(data: bytes) -> bytes:
    """pikepdf空密码解密；pikepdf缺失或解密失败时返回原字节。"""
    try:
        import pikepdf
    except ImportError:
        return data
    try:
        with pikepdf.open(io.BytesIO(data), password="") as pdf:
            out = io.BytesIO()
            pdf.save(out)
            return out.getvalue()
    except Exception:
        return data


# ---------------------------------------------------------------------------
# 坐标定位提取（第7栏车底号 / 第19栏铅封号）
# ---------------------------------------------------------------------------

def _norm_label(text: str) -> str:
    return text.replace(" ", "").replace("\u00a0", "")


def _find_wagon_anchor(words: list[dict]) -> tuple[dict | None, str | None]:
    """定位第7栏"车辆—Вагон"标签词。

    先精确匹配中文+俄文整词；退化时用独立"Вагон"单词并排除同一行右侧
    出现"предоставлен"的（第8栏"Вагон предоставлен"）——这是任务书点名的
    误抓发站代码的干扰源。
    """
    for w in words:
        if _norm_label(w["text"]) == WAGON_ANCHOR_EXACT:
            return w, WAGON_ANCHOR_EXACT
    for w in words:
        if w["text"] == WAGON_ANCHOR_RU:
            same_row_right = [x for x in words
                              if abs(x["top"] - w["top"]) < 3
                              and x["x0"] >= w["x1"]]
            if not any(WAGON_FORBID_NEIGHBOR in x["text"]
                       for x in same_row_right):
                return w, WAGON_ANCHOR_RU
    return None, None


def _find_seal_anchor(words: list[dict]) -> tuple[dict | None, str | None]:
    """定位第19栏下"记号—знаки"子栏标签词（铅封号的取值子框）。"""
    for w in words:
        t = _norm_label(w["text"])
        if t == SEAL_ANCHOR_EXACT or ("记号" in t and t.endswith(SEAL_ANCHOR_RU)):
            return w, t
    for w in words:
        if w["text"] == SEAL_ANCHOR_RU:
            return w, w["text"]
    return None, None


def _vline_xs(page, anchor: dict, depth: float) -> list[float]:
    """锚点下方纵深内出现的竖框线x坐标（去重排序）。"""
    xs = set()
    for e in page.edges:
        if e.get("orientation") != "v":
            continue
        # 竖线需伸入锚点下方取值纵深，排除只在标签行内打转的装饰线
        if e["bottom"] > anchor["bottom"] and e["top"] < anchor["bottom"] + depth:
            xs.add(round(e["x0"], 1))
    return sorted(xs)


def _solve_band(page, anchor: dict, depth: float,
                left_ext: float, right_ext: float) -> tuple[float, float, float, float, bool]:
    """求解锚点标签下方的取值区域 (x0, x1, top, bottom, used_table_lines)。

    优先用表格框线：左右取紧贴锚点的竖线，纵深从锚点底向下延伸depth。
    无框线（纯文本版式）时退化为锚点相对区域，并在调用方note中留痕——
    退化定位的置信度低，批量中占比过高同样是模板漂移信号。
    """
    xs = _vline_xs(page, anchor, depth)
    left = max((x for x in xs if x <= anchor["x0"] + 1), default=None)
    right = min((x for x in xs if x >= anchor["x1"] - 1), default=None)
    if left is not None and right is not None and right > left:
        return left, right, anchor["bottom"] - 1, anchor["bottom"] + depth, True
    return (anchor["x0"] - left_ext, anchor["x1"] + right_ext,
            anchor["bottom"] - 1, anchor["bottom"] + depth, False)


def _words_in_band(words: list[dict], band: tuple) -> list[dict]:
    x0, x1, top, bottom, _ = band
    tol = 2.0
    picked = [w for w in words
              if w["x0"] >= x0 - tol and w["x1"] <= x1 + tol
              and w["top"] >= top - tol and w["bottom"] <= bottom + tol]
    picked.sort(key=lambda w: (w["top"], w["x0"]))
    return picked


def _pick_by_format(candidates: list[str],
                    pattern: re.Pattern) -> tuple[str | None, str, bool]:
    """按格式正则筛选候选值；返回(选中值, 说明, 是否多候选存疑)。"""
    if not candidates:
        return None, "", False
    valid = [c for c in candidates if pattern.fullmatch(c)]
    if not valid:
        return None, "框内候选均不符合格式: " + "/".join(candidates), False
    if len(set(valid)) > 1:
        return valid[0], "框内多个互异候选: " + "/".join(valid), True
    return valid[0], "", False


def _extract_boxed_field(page, words: list[dict], field_name: str,
                         anchor_finder, value_re: re.Pattern,
                         depth: float, left_ext: float,
                         right_ext: float) -> FieldExtraction:
    """通用"锚点+框线"定位提取：车底号与铅封号共用。"""
    anchor, anchor_text = anchor_finder(words)
    if anchor is None:
        return FieldExtraction(
            field=field_name, status=FIELD_ANCHOR_MISSING, anchor=None,
            note="未找到栏位标签锚点词，运单模板可能已变化")

    band = _solve_band(page, anchor, depth, left_ext, right_ext)
    used_lines = band[4]
    inside = _words_in_band(words, band)
    candidates = [w["text"].strip() for w in inside if w["text"].strip()]

    value, fmt_note, ambiguous = _pick_by_format(candidates, value_re)
    note_bits = []
    if not used_lines:
        note_bits.append("无框线，按锚点相对区域定位")
    if fmt_note:
        note_bits.append(fmt_note)
    note = "；".join(note_bits)

    if value is None:
        # 有候选但格式全不符 → 格式异常；框内无任何候选 → 未找到
        status = FIELD_FORMAT_ANOMALY if candidates else FIELD_NOT_FOUND
        return FieldExtraction(
            field=field_name, value=None, status=status,
            candidates=candidates, anchor=anchor_text,
            note=note or "框内无可解析值")

    return FieldExtraction(
        field=field_name, value=value,
        status=FIELD_AMBIGUOUS if ambiguous else FIELD_OK,
        candidates=candidates, anchor=anchor_text,
        note=note or ("框内多个互异候选值，不猜测，请人工确认" if ambiguous else ""))


def _extract_container(page, words: list[dict]) -> FieldExtraction:
    """箱号：既有[A-Z]{4}\\d{7}正则，全文取词；多命中时优先页下部集装箱信息块。

    仅作匹配主键（任务书明确不做真伪核查），方法标记为text正则。
    """
    height = float(page.height)
    seen: list[tuple[str, float]] = []
    for w in words:
        for m in CONTAINER_RE.finditer(w["text"]):
            seen.append((m.group(0), float(w["top"])))
    if not seen:
        return FieldExtraction(field="container_no", status=FIELD_NOT_FOUND,
                               method="text", note="全文未匹配到箱号正则")
    distinct = list(dict.fromkeys(v for v, _ in seen))
    bottom_hits = [v for v, top in seen if top > height * 0.55]
    value = bottom_hits[0] if bottom_hits else distinct[0]
    note = ""
    if len(distinct) > 1:
        note = "多个箱号样式的命中: " + "/".join(distinct) + "；取集装箱信息块（页下部）值"
    return FieldExtraction(field="container_no", value=value,
                           status=FIELD_OK, candidates=distinct,
                           method="text", note=note)


def extract_waybill(data: bytes, filename: str,
                    wagon_re: str = DEFAULT_WAGON_RE,
                    seal_re: str = DEFAULT_SEAL_RE) -> WaybillExtraction:
    """解密+解析一份运单PDF，提取箱号/车底号/铅封号。

    车底号与铅封号严格按坐标/框位定位（见模块docstring的关键约束），
    任何一步不达预期都返回明确的异常状态而非猜测值。
    """
    with open_pdf(data) as pdf:
        page_count = len(pdf.pages)
        page = pdf.pages[0]
        words = page.extract_words(use_text_flow=False, keep_blank_chars=False)

        container = _extract_container(page, words)
        wagon = _extract_boxed_field(
            page, words, "wagon_no", _find_wagon_anchor,
            re.compile(wagon_re), WAGON_BAND_DEPTH, WAGON_BAND_LEFT, WAGON_BAND_RIGHT)
        seal = _extract_boxed_field(
            page, words, "seal_no", _find_seal_anchor,
            re.compile(seal_re), SEAL_BAND_DEPTH, SEAL_BAND_LEFT, SEAL_BAND_RIGHT)

    return WaybillExtraction(filename=filename, container=container,
                             wagon_no=wagon, seal_no=seal,
                             page_count=page_count)


def extract_waybill_safe(data: bytes, filename: str,
                         **kwargs) -> WaybillExtraction:
    """extract_waybill 的批量友好包装：单份文件解析崩溃不拖垮整批核对，
    转为 read_error 异常结果进入批级异常统计（不静默跳过）。"""
    try:
        return extract_waybill(data, filename, **kwargs)
    except Exception as exc:  # noqa: BLE001  单份坏文件按异常结果计入，不中断批次
        note = f"文件无法解析（损坏或不支持的加密）：{type(exc).__name__}: {exc}"
        return WaybillExtraction(
            filename=filename,
            container=FieldExtraction(field="container_no",
                                      status=FIELD_READ_ERROR, note=note),
            wagon_no=FieldExtraction(field="wagon_no",
                                     status=FIELD_READ_ERROR, note=note),
            seal_no=FieldExtraction(field="seal_no",
                                    status=FIELD_READ_ERROR, note=note))


# ---------------------------------------------------------------------------
# 现场装车清单Excel解析
# ---------------------------------------------------------------------------

#: 清单表头列名识别（包含式匹配；兼容"铅封号/封号"两种写法）
_SEQ_HEADER = ("序号",)
_TYPE_HEADER = ("车种", "车板种类")
_WAGON_HEADER = ("车底号", "车板号")
_CONTAINER_HEADER = ("箱号",)
_SEAL_HEADER = ("铅封号", "封号")


def _find_header_row(ws) -> tuple[int, dict[str, int]] | None:
    """在前若干行里找同时含"箱号"与"铅封号/封号"的表头行，返回(行号, 列映射)。"""
    for row in ws.iter_rows(min_row=1, max_row=20):
        texts = [str(c.value).strip() if c.value is not None else ""
                 for c in row]
        joined = [t.replace(" ", "") for t in texts]
        if not any("箱号" in t for t in joined):
            continue
        if not any(any(h in t for h in _SEAL_HEADER) for t in joined):
            continue
        colmap: dict[str, int] = {}
        for idx, t in enumerate(joined, 1):
            if not t:
                continue
            if "序号" in t and "seq" not in colmap:
                colmap["seq"] = idx
            elif any(h in t for h in _WAGON_HEADER) and "wagon" not in colmap:
                colmap["wagon"] = idx
            elif "车种" in t and "type" not in colmap:
                colmap["type"] = idx
            elif "箱号" in t and "container" not in colmap:
                colmap["container"] = idx
            elif any(h in t for h in _SEAL_HEADER) and "seal" not in colmap:
                colmap["seal"] = idx
        if "container" in colmap and "seal" in colmap:
            return row[0].row, colmap
    return None


def _cell_text(value) -> str | None:
    """单元格归一化：数字去尾零（260281.0→"260281"），文本去空白。"""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    return text or None


def _parse_rows(ws, header_row: int, colmap: dict[str, int]) -> list[PackingRow]:
    rows: list[PackingRow] = []
    for row in ws.iter_rows(min_row=header_row + 1):
        cells = {k: (row[idx - 1].value if idx <= len(row) else None)
                 for k, idx in colmap.items()}
        container = _cell_text(cells.get("container"))
        if not container:
            continue  # 空行/备注行
        rows.append(PackingRow(
            row_no=row[0].row,
            seq=_cell_text(cells.get("seq")),
            wagon_type=_cell_text(cells.get("type")),
            wagon_no=_cell_text(cells.get("wagon")),
            container_no=container.strip().upper().replace(" ", ""),
            seal_no=_cell_text(cells.get("seal")),
        ))
    return rows


def load_packing_list(source) -> list[PackingRow]:
    """解析A同事现场装车清单Excel（发运清单电子版口径，表头自动识别）。

    source：文件路径或已读入的bytes（网页上传场景）。
    工作表选择：真实清单文件常同时含空白打印模板与电子版数据表，多个工作表
    都可能带"箱号/封号"表头——逐表解析后取有效箱号行（符合箱号正则）最多的
    一张，避免选到空模板；保留Excel实际行号与"序号"列值，便于核对界面用
    业务口径指行（任务书"现场清单第48行"即序号列口径）。
    """
    import openpyxl

    if isinstance(source, (bytes, bytearray)):
        wb = openpyxl.load_workbook(io.BytesIO(source), data_only=True)
    else:
        wb = openpyxl.load_workbook(source, data_only=True)

    best: tuple[int, list[PackingRow]] | None = None
    for ws in wb.worksheets:
        found = _find_header_row(ws)
        if not found:
            continue
        header_row, colmap = found
        rows = _parse_rows(ws, header_row, colmap)
        valid = sum(1 for r in rows if CONTAINER_RE.fullmatch(r.container_no or ""))
        if best is None or valid > best[0]:
            best = (valid, rows)
    if not best or best[0] == 0:
        raise ValueError("未在任一工作表找到含有效箱号数据的清单"
                         "（表头需同时含'箱号'与'铅封号/封号'列）")
    return best[1]


# ---------------------------------------------------------------------------
# 批量核对（箱号主键匹配 + 车底号/铅封号比较 + 批级合理性告警）
# ---------------------------------------------------------------------------

@dataclass
class CompareConfig:
    """核对口径配置；正则与提取保持同一来源，改这里即全局生效。"""
    wagon_re: str = DEFAULT_WAGON_RE
    seal_re: str = DEFAULT_SEAL_RE
    anomaly_ratio_threshold: float = field(
        default_factory=default_anomaly_ratio)


def _digits(text: str | None) -> str:
    return re.sub(r"\D", "", text or "")


def _compare_pair(ext: WaybillExtraction, prow: PackingRow | None) -> dict:
    """单份PDF与清单行的核对明细（字段级，异常与不一致分别留痕）。"""
    field_results = {}
    for fname, pdf_field, excel_value in (
            ("wagon_no", ext.wagon_no, prow.wagon_no if prow else None),
            ("seal_no", ext.seal_no, prow.seal_no if prow else None)):
        entry: dict = {"pdf_value": pdf_field.value,
                       "pdf_status": pdf_field.status,
                       "excel_value": excel_value,
                       "result": "not_compared", "note": ""}
        if pdf_field.is_anomaly:
            entry["note"] = f"PDF提取异常（{pdf_field.status}）"
        elif not excel_value:
            entry["note"] = "清单该字段为空，未比较"
        elif _digits(pdf_field.value) == _digits(excel_value):
            entry["result"] = "match"
        else:
            entry["result"] = "mismatch"
        field_results[fname] = entry

    mismatched = [f for f, e in field_results.items() if e["result"] == "mismatch"]
    anomalous = [f for f, e in field_results.items() if e["pdf_status"] != FIELD_OK]

    if prow is None:
        status = ST_PDF_UNMATCHED
    elif anomalous:
        status = ST_EXTRACTION_ANOMALY
    elif mismatched:
        status = {("wagon_no",): ST_WAGON_MISMATCH,
                  ("seal_no",): ST_SEAL_MISMATCH}.get(tuple(mismatched),
                                                      ST_BOTH_MISMATCH)
    elif any(not e["excel_value"] for e in field_results.values()):
        status = ST_DATA_MISSING
    else:
        status = ST_MATCH
    return {"status": status, "mismatch_fields": mismatched,
            "anomaly_fields": anomalous, "field_results": field_results}


#: "未匹配"提示语必须说明两种可能（任务书：不能简单写"缺失"误导方向）
PDF_UNMATCHED_TEXT = ("该PDF运单的箱号在A同事现场清单中未找到——可能是运单里箱号"
                      "本身录入有误，也可能是清单里没有这条记录，请人工判断")
EXCEL_UNMATCHED_TEXT = ("现场清单有此箱号，但上传的PDF运单中没有匹配到——可能是漏传"
                        "了这份运单，也可能是PDF中箱号录错（该PDF会出现在'PDF未匹配'区）")


def compare_batch(rows: list[PackingRow],
                  extractions: list[WaybillExtraction],
                  config: CompareConfig | None = None) -> dict:
    """以箱号为主键做批量核对，并执行批级合理性校验（防静默失效）。

    返回：
      results     每份PDF/每条清单一行的核对结果（不一致的排最前）
      summary     各状态计数
      anomaly     批级异常统计与告警（ratio超阈值→batch_warning=True）
    """
    cfg = config or CompareConfig()

    by_container: dict[str, list[PackingRow]] = {}
    for r in rows:
        if r.container_no:
            by_container.setdefault(r.container_no.strip().upper(), []).append(r)

    results: list[dict] = []
    matched_excel_rows: set[int] = set()

    for ext in extractions:
        item = {
            "filename": ext.filename,
            "pdf_container": ext.container.value,
            "seq": None, "row_no": None, "wagon_type": None,
            "excel_wagon": None, "excel_seal": None,
            "notes": [f.note for f in ext.fields() if f.note],
        }
        key = (ext.container.value or "").strip().upper()
        container_bad = ext.container.status != FIELD_OK
        prow = None
        if not container_bad and key in by_container:
            prow = by_container[key][0]
            item.update(seq=prow.seq, row_no=prow.row_no,
                        wagon_type=prow.wagon_type,
                        excel_wagon=prow.wagon_no, excel_seal=prow.seal_no)
            matched_excel_rows.add(prow.row_no)
            if len(by_container[key]) > 1:
                item["notes"].append(
                    f"清单中箱号{key}出现{len(by_container[key])}行，按第一行核对")
        item.update(_compare_pair(ext, prow))
        if container_bad:
            # 箱号没提到/文件不可读：谈不上"未匹配"（那会误导成箱号录错），
            # 归入提取异常（任务书：不静默输出看似正常的结果）
            item["status"] = ST_EXTRACTION_ANOMALY
        results.append(item)

    # 清单侧：有箱号但没等到PDF的行
    for r in rows:
        if r.row_no not in matched_excel_rows:
            results.append({
                "filename": None, "pdf_container": None,
                "seq": r.seq, "row_no": r.row_no, "wagon_type": r.wagon_type,
                "excel_container": r.container_no,
                "excel_wagon": r.wagon_no, "excel_seal": r.seal_no,
                "status": ST_EXCEL_UNMATCHED, "mismatch_fields": [],
                "anomaly_fields": [], "field_results": {},
                "notes": [EXCEL_UNMATCHED_TEXT],
            })

    # 展示顺序：两项不一致 > 车底号不一致 > 铅封号不一致 > 其余按输入序
    _severity = {ST_BOTH_MISMATCH: 0, ST_WAGON_MISMATCH: 1, ST_SEAL_MISMATCH: 2}
    results.sort(key=lambda x: (_severity.get(x["status"], 3),))

    summary = {
        "total_pdfs": len(extractions),
        "total_excel_rows": len(rows),
    }
    for st in (ST_MATCH, ST_WAGON_MISMATCH, ST_SEAL_MISMATCH, ST_BOTH_MISMATCH,
               ST_PDF_UNMATCHED, ST_EXCEL_UNMATCHED, ST_EXTRACTION_ANOMALY,
               ST_DATA_MISSING):
        summary[st] = sum(1 for r in results if r["status"] == st)

    anomaly = _batch_anomaly_stats(extractions, cfg.anomaly_ratio_threshold)
    return {"results": results, "summary": summary, "anomaly": anomaly,
            "config": {"wagon_re": cfg.wagon_re, "seal_re": cfg.seal_re,
                       "anomaly_ratio_threshold": cfg.anomaly_ratio_threshold}}


def _batch_anomaly_stats(extractions: list[WaybillExtraction],
                         threshold: float) -> dict:
    """批级合理性校验：异常占比超阈值→告警（任务书"防止静默出错"）。

    区分两类信号：
      anchor_missing 类 —— 模板漂移的直接证据，文案指向"模板可能已变化"；
      其他异常        —— 单份文件问题，占比过高时同样升级为模板告警。
    """
    total = len(extractions)
    bad = [e for e in extractions if e.anomaly_fields()]
    anchor_missing = [e for e in bad
                      if any(f.status == FIELD_ANCHOR_MISSING
                             for f in e.fields())]
    ratio = (len(bad) / total) if total else 0.0
    warning = bool(total) and ratio > threshold

    if anchor_missing:
        reason = (f"{len(anchor_missing)}/{total} 份运单找不到栏位锚点词"
                  f"（第7栏车辆/第19栏封印），模板结构可能已变化")
    elif warning:
        reason = f"{len(bad)}/{total} 份运单提取异常，占比 {ratio:.0%} 超过阈值"
    else:
        reason = ""

    return {
        "total": total,
        "anomaly_files": [e.filename for e in bad],
        "anomaly_ratio": round(ratio, 4),
        "threshold": threshold,
        "batch_warning": warning,
        "anchor_missing_files": [e.filename for e in anchor_missing],
        "warning_text": (
            f"⚠️ 运单模板可能已变化，建议人工抽查：{reason}。"
            f"本系统按确定性规则（正则+坐标定位）提取，模板变化会导致"
            f"规则失效，已停止给出'看似正常'的自动结论。" if warning else ""),
    }
