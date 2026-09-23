# -*- coding: utf-8 -*-
"""数据核对模块 Excel 导入（任务书 §五）。

支持两类导入（模板字段命名与任务书数据结构一致）：
  kind="trip"    班列基础信息 + 联运费用结算明细 + 实付
  kind="subsidy" 补贴测算（三方补贴列 + 预测差额 + 原因分析）

设计口径：
  - 兼容真实台账的常见写法：双行表头（分组行/子列行）、`2025.7.18`、
    `2025/10/10临`、Excel日期序列号；**日期单元格后缀 临→L、图→T**，
    这正是同日临/图两趟班列拆成两个编号的入口（无后缀默认 T）；
  - 表头识别按"归一化别名表"匹配（去空格/单位/括号/百分号后比对），
    不依赖列序，兼容 开行日期≈发运日期、到站≈目的地、合计≈结算合计 等差异；
  - 发站/口岸/目的地允许填中文名或缩写，经代码字典双向映射（未登记→报错提示）；
  - 预览（detect）与应用（apply）分离：预览只读并标注 新增/一致/冲突，
    应用时逐行携带人工决定（create/overwrite/keep），绝不静默覆盖；
    应用端二次校验（并发下编号已存在→按冲突跳过而非重复建）。
"""

from __future__ import annotations

import io
import re
from datetime import date, datetime, timedelta

from openpyxl import Workbook, load_workbook

import train_number
import train_recon
import train_store
from train_number import TrainNumberError
from train_store import TrainStoreError

IMPORT_KINDS = ("trip", "subsidy")

KIND_LABELS = {"trip": "班列结算导入", "subsidy": "补贴测算导入"}

# 模板/识别共用的列结构：(字段键, 分组行, 子列行)
TRIP_COLUMNS = [
    ("dep_date", "班列基础信息", "发运日期"),
    ("station", "班列基础信息", "发站"),
    ("port", "班列基础信息", "口岸"),
    ("dest", "班列基础信息", "目的地"),
    ("train_type", "班列基础信息", "类型(L/T)"),
    ("goods_name", "班列基础信息", "货物品名"),
    ("wagon_count", "班列基础信息", "车数"),
    ("container_40hd", "柜量（40HD/20HD/折合TEU）", "40HD"),
    ("container_20hd", "柜量（40HD/20HD/折合TEU）", "20HD"),
    ("teu_total", "柜量（40HD/20HD/折合TEU）", "折合TEU"),
    ("rail_freight", "结算明细（元）", "铁路运费"),
    ("customs_fee", "结算明细（元）", "报关费"),
    ("service_fee", "结算明细（元）", "服务费"),
    ("other_fee", "结算明细（元）", "其他费用"),
    ("settle_total", "结算明细（元）", "合计"),
    ("actual_freight", "实付（元）", "运费"),
    ("actual_paid_at", "实付（元）", "实付时间"),
    ("diff_reason", "备注", "差异原因"),
]

SUBSIDY_COLUMNS = [
    ("dep_date", "班列信息", "发运日期"),
    ("train_type", "班列信息", "类型(L/T)"),
    ("station", "班列信息", "发站"),
    ("port", "班列信息", "口岸"),
    ("dest", "班列信息", "目的地"),
    ("dt_supply_100", "大同供应链补贴测算", "补贴资料100%"),
    ("dt_supply_70", "大同供应链补贴测算", "补贴70%"),
    ("ly_advance_100", "联运公司补贴测算", "联运垫付补贴100%"),
    ("ly_recover_70", "联运公司补贴测算", "需要给联运回款70%"),
    ("auth_confirm_100", "上级拨付", "确认补贴100%"),
    ("auth_advance_70", "上级拨付", "预拨付70%"),
    ("auth_remain_30", "上级拨付", "剩余30%"),
    ("forecast_diff", "预测", "预测补贴差额"),
    ("diff_reason", "预测", "原因分析"),
]

# 表头归一化别名表（归一化后比对；归一化=仅保留小写字母数字与汉字）
HEADER_ALIASES = {
    "dep_date": ["发运日期", "开行日期", "发运时间", "班列发运日期", "日期"],
    "station": ["发站", "发运站", "始发站"],
    "port": ["口岸", "出境口岸", "口岸站"],
    "dest": ["目的地", "目的国", "到站", "目的国地区", "目的国地区缩写"],
    "train_type": ["类型lt", "类型", "车次性质", "lt"],
    "goods_name": ["货物品名", "品名"],
    "wagon_count": ["车数", "车皮数"],
    "container_40hd": ["柜量40hd", "40hd", "40柜", "大柜"],
    "container_20hd": ["柜量20hd", "20hd", "20柜", "小柜"],
    "teu_total": ["柜量折合teu", "折合teu", "teu"],
    "rail_freight": ["铁路运费", "运费"],
    "customs_fee": ["报关费"],
    "service_fee": ["服务费", "代理费"],
    "other_fee": ["其他费用", "其他"],
    "settle_total": ["结算合计", "结算金额合计", "合计", "结算金额"],
    "actual_freight": ["实付运费", "实付元运费", "实付"],
    "actual_paid_at": ["实付时间", "付款日期", "实付日期"],
    "diff_reason": ["差异原因", "原因分析", "原因"],
    "dt_supply_100": ["大同供应链补贴测算补贴资料100", "补贴资料100", "大同供应链补贴100"],
    "dt_supply_70": ["大同供应链补贴测算补贴70", "补贴70"],
    "ly_advance_100": ["联运公司补贴测算联运垫付补贴100", "联运垫付补贴100"],
    "ly_recover_70": ["联运公司补贴测算需要给联运回款70", "需要给联运回款70", "需回款给联运的70"],
    "auth_confirm_100": ["上级拨付确认补贴100", "确认补贴100"],
    "auth_advance_70": ["上级拨付预拨付70", "预拨付70"],
    "auth_remain_30": ["上级拨付剩余30", "上级拨付30", "剩余30"],
    "forecast_diff": ["预测补贴差额", "预测补贴差额及原因分析"],
}

# 归一化：仅保留 小写字母/数字/汉字（括号、单位"元"、百分号、斜杠等全部剔除）
_NORM_RE = re.compile(r"[^a-z0-9\u4e00-\u9fff]+")
# 日期单元格中的日期主干 + 临/图 后缀
_DATE_RE = re.compile(r"(\d{4})\s*[./\-年]\s*(\d{1,2})\s*[./\-月]\s*(\d{1,2})")
_TYPE_SUFFIX_RE = re.compile(r"[临图]")
_EXCEL_EPOCH = date(1899, 12, 30)


class ImportError_(ValueError):
    """导入解析错误（文件格式、表头缺失等）。"""


def _norm_header(text) -> str:
    return _NORM_RE.sub("", str(text or "").lower())


def _match_header(text) -> str | None:
    """表头文本 → 字段键（按别名表）。"""
    norm = _norm_header(text)
    if not norm:
        return None
    for field, aliases in HEADER_ALIASES.items():
        if norm in aliases:
            return field
    return None


def parse_dep_cell(value) -> tuple[date | None, str | None]:
    """发运日期单元格 → (date, 类型提示 L/T/None)。

    支持 datetime 单元格、Excel 序列号、2025.7.18 / 2025/10/10临 /
    2025-10-10 / 2025年10月10日(图) 等写法；临→L、图→T（任务书§五，
    同日临/图拆两趟的关键入口）。解析失败返回 (None, None)。
    """
    if value is None:
        return None, None
    if isinstance(value, datetime):
        return value.date(), None
    if isinstance(value, date):
        return value, None
    if isinstance(value, (int, float)):        # Excel 日期序列号
        if 0 < float(value) < 100000:
            return _EXCEL_EPOCH + timedelta(days=int(value)), None
        return None, None
    text = str(value).strip()
    if not text:
        return None, None
    type_hint = None
    m = _TYPE_SUFFIX_RE.search(text)
    if m:
        type_hint = train_number.normalize_train_type(m.group())
    dm = _DATE_RE.search(text)
    if not dm:
        return (None, type_hint) if type_hint else (None, None)
    try:
        return date(int(dm.group(1)), int(dm.group(2)), int(dm.group(3))), type_hint
    except ValueError:
        return None, type_hint


def _name_or_code(value: str, category: str,
                  name_map: dict[str, str]) -> str | None:
    """单元格值 → 缩写：已是登记缩写直接用；中文名经字典反查；否则 None。"""
    text = str(value or "").strip()
    if not text:
        return None
    code = text.upper()
    codes = name_map.get(category) or {}
    if code in codes:
        return code
    for c, name in codes.items():
        if text == name or text in name or name in text:
            return c
    return None


def _locate_headers(ws, expected_fields: list[str]):
    """在前6行内定位表头（支持 双行表头/单行表头），返回 (field→col, 数据起始行)。"""
    rows = list(ws.iter_rows(min_row=1, max_row=6, values_only=True))
    width = max((len(r) for r in rows), default=0)
    best = None   # (matched_count, mapping, data_start)
    for i in range(len(rows)):
        single = {}
        for col, cell in enumerate(rows[i]):
            f = _match_header(cell)
            if f:
                single.setdefault(col, f)
        if single:
            count = len(set(single.values()) & set(expected_fields))
            if best is None or count > best[0]:
                best = (count, single, i + 1)
        if i + 1 < len(rows):
            combined = {}
            for col in range(width):
                g = str(rows[i][col] or "").strip()
                s = str(rows[i + 1][col] or "").strip()
                if not g and not s:
                    continue
                # 优先级：分组+子列拼接（"实付(元)"/"运费"→实付运费）
                # → 子列单独（"柜量(…)"/"40HD"→40HD）→ 分组单独（合并单元格）
                f = _match_header(g + s) or _match_header(s) or _match_header(g)
                if f:
                    combined.setdefault(col, f)
            if combined:
                count = len(set(combined.values()) & set(expected_fields))
                if best is None or count > best[0]:
                    # rows[i]/rows[i+1] 为双行表头（0基）→ 数据从表头下一行开始
                    best = (count, combined, i + 3)
    if not best or best[0] < 2:
        raise ImportError_(
            "未识别到有效的表头（至少需要匹配2个已知列名）。请下载模板按格式填写，"
            "或参考真实台账的列名（发运日期/发站/口岸/目的地…）。")
    return best[1], best[2]


def _cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_sheet(kind: str, data: bytes) -> list[dict]:
    """解析上传的 xlsx → 归一化行列表（每行含 dep_date/train_type/组件/数值字段）。"""
    if kind not in IMPORT_KINDS:
        raise ImportError_(f"导入类别不合法：{kind!r}")
    columns = TRIP_COLUMNS if kind == "trip" else SUBSIDY_COLUMNS
    expected = [c[0] for c in columns]
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportError_(f"无法读取 Excel 文件：{exc}")
    ws = wb.active
    mapping, data_start = _locate_headers(ws, expected)
    rows_out: list[dict] = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=data_start, values_only=True),
                                  start=data_start):
        values: dict = {}
        for col, field in mapping.items():
            if col < len(row):
                values[field] = row[col]
        dep_date, type_hint = parse_dep_cell(values.get("dep_date"))
        if dep_date is None:
            # 无日期的行：全空=空行跳过；有内容=坏行报错（如合计行/#REF!行）
            if any(_cell_str(v) for v in row or []):
                rows_out.append({"row_index": row_idx, "error":
                                 "发运日期无法解析（跳过该行，请修正后单独补录）"})
            continue
        rec = {"row_index": row_idx, "dep_date": dep_date.isoformat()}
        try:
            rec["train_type"] = (train_number.normalize_train_type(
                values.get("train_type")) if _cell_str(values.get("train_type"))
                else (type_hint or train_number.TYPE_SCHEDULED))
        except TrainNumberError:
            rec["train_type"] = type_hint or train_number.TYPE_SCHEDULED
        for field, _group, _sub in columns:
            if field in ("dep_date", "train_type"):
                continue
            rec[field] = _cell_str(values.get(field))
        rows_out.append(rec)
    wb.close()
    return rows_out


def _compare_numbers(new: str, old) -> bool:
    """数值等价比较（'2666062.00' == 2666062.0）。无法解析时按字符串比较。"""
    a, b = train_recon.to_decimal(new), train_recon.to_decimal(old)
    if a is None or b is None:
        return _cell_str(new) == _cell_str(old)
    return a == b


# 与既有记录比对的字段（trip=基础+结算；subsidy=全部补贴数值+原因）
_TRIP_COMPARE = ["goods_name", "wagon_count", "container_40hd", "container_20hd",
                 "teu_total", "rail_freight", "customs_fee", "service_fee",
                 "other_fee", "settle_total", "actual_freight"]
_SUBSIDY_COMPARE = [c for c in train_store.SUBSIDY_AMOUNT_COLS] + ["diff_reason"]


def detect(kind: str, data: bytes) -> dict:
    """解析 + 与既有记录比对 → 差异预览（绝不写库）。

    每行 action：new（新增）/ same（与现有一致，无需导入）/
    conflict（与现有不一致，需人工决定 覆盖/保留）/ error（无法落库的原因）。
    """
    raw_rows = parse_sheet(kind, data)
    code_map = train_store.load_code_map()
    name_maps = {
        "station": train_store.load_code_map().get("station", {}),
        "port": code_map.get("port", {}),
        "dest": code_map.get("dest", {}),
    }
    rows: list[dict] = []
    stats = {"new": 0, "same": 0, "conflict": 0, "error": 0}
    for rec in raw_rows:
        if rec.get("error"):                    # 解析失败的坏行（如合计行/#REF!行）
            rows.append({"row_index": rec["row_index"], "dep_date": "",
                         "train_type": "", "trip_no": "", "action": "error",
                         "message": rec["error"], "values": {}, "diffs": {}})
            stats["error"] += 1
            continue
        out = {"row_index": rec["row_index"],
               "dep_date": rec["dep_date"], "train_type": rec["train_type"],
               "values": {k: v for k, v in rec.items()
                          if k not in ("row_index",)},}
        # 中文名 → 缩写（未登记 → error，提示先在代码字典登记）
        codes = {}
        for cat in ("station", "port", "dest"):
            code = _name_or_code(rec.get(cat), cat, name_maps)
            if not code:
                out.update(action="error", trip_no="",
                           message=f"第{rec['row_index']}行：{ {'station':'发站','port':'口岸','dest':'目的地'}[cat] }"
                                   f"“{rec.get(cat)}”未在代码字典登记，请先登记后再导入")
                break
            codes[cat] = code
        else:
            base = train_number.build_base(rec["dep_date"], codes["station"],
                                           codes["port"], codes["dest"],
                                           rec["train_type"], code_map)
            existing = train_store.list_trips(
                date_from=rec["dep_date"], date_to=rec["dep_date"],
                station=codes["station"], port=codes["port"],
                dest=codes["dest"])
            existing = [t for t in existing if t["train_type"] == rec["train_type"]]
            out["values"].update(codes)
            if not existing:
                out.update(action="new", trip_no=base, message="将新建班列并写入数据")
                stats["new"] += 1
            else:
                trip = min(existing, key=lambda t: t["suffix"])
                out["trip_no"] = trip["trip_no"]
                compare = _TRIP_COMPARE if kind == "trip" else _SUBSIDY_COMPARE
                # 比对池：trip=主档+结算明细；subsidy=补贴行
                old_pool = ({**trip, **(train_store.get_settlement(trip["trip_no"]) or {})}
                            if kind == "trip"
                            else (train_store.get_subsidy(trip["trip_no"]) or {}))
                diffs = {}
                for field in compare:
                    new_v = rec.get(field)
                    if new_v is None or _cell_str(new_v) == "":
                        continue           # Excel 未填 ≠ 冲突，保留现状
                    old_v = old_pool.get(field)
                    same = _compare_numbers(new_v, old_v) if field != "diff_reason" \
                        else _cell_str(new_v) == _cell_str(old_v)
                    if not same:
                        diffs[field] = {"old": "" if old_v is None else str(old_v),
                                        "new": str(new_v)}
                if not diffs:
                    out.update(action="same", diffs={},
                               message="与现有记录一致，无需导入")
                    stats["same"] += 1
                else:
                    out.update(action="conflict", diffs=diffs,
                               message="与现有记录不一致，请选择 覆盖 或 保留")
                    stats["conflict"] += 1
        rows.append(out)
    return {"kind": kind, "rows": rows, "stats": stats,
            "total": len(rows),
            "message": (f"共 {len(rows)} 行：新增 {stats['new']}，一致 {stats['same']}，"
                        f"冲突 {stats['conflict']}，错误 {stats['error']}")}


def apply(kind: str, decisions: list[dict], by: str) -> dict:
    """按人工决定应用导入（decision: create/overwrite/keep）。

    应用端二次校验：create 时编号已存在（并发）→ 按 conflict 跳过；
    subsidy 的班列不存在 → error。返回逐行结果供审计与界面提示。
    """
    code_map = train_store.load_code_map()
    name_maps = {"station": code_map.get("station", {}),
                 "port": code_map.get("port", {}), "dest": code_map.get("dest", {})}
    results = []
    counts = {"created": 0, "updated": 0, "kept": 0, "error": 0}
    for item in decisions:
        decision = item.get("decision")
        values = dict(item.get("values") or {})
        try:
            codes = {}
            for cat in ("station", "port", "dest"):
                code = _name_or_code(values.get(cat), cat, name_maps)
                if not code:
                    raise ImportError_(
                        f"{ {'station':'发站','port':'口岸','dest':'目的地'}[cat] }"
                        f"“{values.get(cat)}”未登记")
                codes[cat] = code
            values.update(codes)
            base = train_number.build_base(values["dep_date"], codes["station"],
                                           codes["port"], codes["dest"],
                                           values.get("train_type", "T"), code_map)
            existing = [t for t in train_store.list_trips(
                date_from=values["dep_date"], date_to=values["dep_date"],
                station=codes["station"], port=codes["port"], dest=codes["dest"])
                if t["train_type"] == values.get("train_type", "T")]
            if decision == "keep":
                counts["kept"] += 1
                results.append({"row_index": item.get("row_index"),
                                "status": "kept",
                                "trip_no": existing[0]["trip_no"] if existing else "",
                                "message": "按人工选择保留现有记录"})
                continue
            if kind == "trip":
                if not existing:
                    basic = {k: values.get(k) for k in
                             ("goods_name", "wagon_count", "container_40hd",
                              "container_20hd", "teu_total")}
                    trip = train_store.create_trip(
                        values["dep_date"], codes["station"], codes["port"],
                        codes["dest"], values.get("train_type", "T"), by, basic)
                    settle_fields = {k: values[k] for k in
                                     ("rail_freight", "customs_fee", "service_fee",
                                      "other_fee", "settle_total", "actual_freight")
                                     if _cell_str(values.get(k))}
                    if settle_fields:
                        train_store.upsert_settlement(trip["trip_no"],
                                                      settle_fields, by)
                    counts["created"] += 1
                    results.append({"row_index": item.get("row_index"),
                                    "status": "created",
                                    "trip_no": trip["trip_no"],
                                    "message": "已新建班列并写入结算数据"})
                elif decision == "overwrite":
                    trip_no = min(existing, key=lambda t: t["suffix"])["trip_no"]
                    basic = {k: values[k] for k in
                             ("goods_name", "wagon_count", "container_40hd",
                              "container_20hd", "teu_total")
                             if _cell_str(values.get(k))}
                    if basic:
                        train_store.update_trip_basic(trip_no, basic, by)
                    settle_fields = {k: values[k] for k in
                                     ("rail_freight", "customs_fee", "service_fee",
                                      "other_fee", "settle_total", "actual_freight",
                                      "diff_reason")
                                     if _cell_str(values.get(k))}
                    if settle_fields:
                        train_store.upsert_settlement(trip_no, settle_fields, by)
                    counts["updated"] += 1
                    results.append({"row_index": item.get("row_index"),
                                    "status": "updated", "trip_no": trip_no,
                                    "message": "已按人工选择覆盖现有记录"})
                else:
                    counts["kept"] += 1
                    results.append({"row_index": item.get("row_index"),
                                    "status": "kept", "trip_no": existing[0]["trip_no"],
                                    "message": "班列已存在且未选择覆盖，跳过"})
            else:   # kind == "subsidy"：必须挂到已存在班列
                if not existing:
                    raise ImportError_(
                        f"班列 {base} 尚未登记，请先通过“班列结算导入”或手工登记班列")
                trip_no = min(existing, key=lambda t: t["suffix"])["trip_no"]
                if decision != "overwrite":
                    counts["kept"] += 1
                    results.append({"row_index": item.get("row_index"),
                                    "status": "kept", "trip_no": trip_no,
                                    "message": "按人工选择保留现有记录"})
                    continue
                sub_fields = {k: values[k] for k in train_store.SUBSIDY_AMOUNT_COLS
                              if _cell_str(values.get(k))}
                if "diff_reason" in values and _cell_str(values.get("diff_reason")):
                    sub_fields["diff_reason"] = values["diff_reason"]
                train_store.upsert_subsidy(trip_no, sub_fields, by)
                counts["updated"] += 1
                results.append({"row_index": item.get("row_index"),
                                "status": "updated", "trip_no": trip_no,
                                "message": "已按人工选择覆盖补贴数据"})
        except (TrainNumberError, TrainStoreError, ImportError_) as exc:
            counts["error"] += 1
            results.append({"row_index": item.get("row_index"), "status": "error",
                            "trip_no": "", "message": str(exc)})
    # 导入后刷新重复值异常快照（补贴类或写入过数据的）
    if counts["created"] or counts["updated"]:
        train_store.refresh_anomalies(by)
    return {"kind": kind, "counts": counts, "results": results,
            "message": (f"处理完成：新建 {counts['created']}，覆盖 {counts['updated']}，"
                        f"保留 {counts['kept']}，错误 {counts['error']}")}


# ---------------------------------------------------------------- 模板生成

def build_template(kind: str) -> bytes:
    """生成导入模板 xlsx（双行表头 + 一行脱敏示例数据）。"""
    if kind not in IMPORT_KINDS:
        raise ImportError_(f"导入类别不合法：{kind!r}")
    columns = TRIP_COLUMNS if kind == "trip" else SUBSIDY_COLUMNS
    wb = Workbook()
    ws = wb.active
    ws.title = KIND_LABELS[kind]
    ws.append([c[1] for c in columns])       # 分组行
    ws.append([c[2] for c in columns])       # 子列行
    if kind == "trip":
        ws.append(["2025-07-18", "平旺", "满洲里", "俄罗斯", "T", "汽车零配件",
                   50, 25, 50, 100, 2000000.00, 50000.00, 30000.00, 10000.00,
                   2090000.00, 2090000.00, "2025-08-05", ""])
    else:
        ws.append(["2025-07-18", "T", "平旺", "满洲里", "俄罗斯",
                   1000000.00, 700000.00, 1000000.00, 700000.00,
                   1000000.00, 700000.00, 300000.00, 0.00, "示例数据（脱敏）"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
